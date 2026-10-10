from pathlib import Path
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import io
import fitz
import pytesseract
from PIL import Image
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

class DocumentUtils:
    _instance = None
    def __new__(cls, *args, **kwargs):
        if cls._instance is None:
            cls._instance = super().__new__(cls)
        return cls._instance
    
    def __init__(self):
        if hasattr(self, '_initialized'):
            return
        else:
            self._initialized = True
            self.extracted_text = ""

    def extract_text(self, file_path: Path) -> str:
        doc_type = file_path.suffix.lower()

        if doc_type == ".pdf":
            self.extracted_text = self.extract_from_pdf(file_path)
        elif doc_type == ".docx":
            self.extracted_text = self.extract_from_docx(file_path)
        elif doc_type == ".pptx":
            self.extracted_text = self.extract_from_pptx(file_path)
        elif doc_type == ".txt":
            self.extracted_text = self.extract_from_txt(file_path)
        else:
            raise ValueError(f"Unsupported document type: {doc_type}")

        return self.extracted_text


    @staticmethod
    def extract_from_pdf(file_path) -> str:
        reader = PdfReader(file_path)
        doc = fitz.open(str(file_path))
        try:
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = (page.extract_text() or "").strip()
                if len(text) < 20:
                    pix = doc[i].get_pixmap(dpi=200)
                    image = Image.open(io.BytesIO(pix.tobytes("png")))
                    text = pytesseract.image_to_string(image).strip()
                if text:
                    pages_text.append(text)
            return "\n\n".join(pages_text)
        finally:
            doc.close()

        

    @staticmethod
    def extract_from_docx(file_path: Path) -> str:
        file = Document(file_path)
        full_text = ""
        for paragraph in file.paragraphs:
            text = paragraph.text
            if text:
                full_text += text + "\n"
        return full_text

    @staticmethod
    def extract_from_pptx(file_path: Path) -> str:
        file = Presentation(file_path)
        extracted_text = []

        for slide in file.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text.strip():
                                extracted_text.append(run.text)

        return " ".join(extracted_text)


    @staticmethod
    def extract_from_txt(file_path: Path) -> str:

        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()

    @staticmethod
    def chunk_text(text: str, chunksize: int = 500, overlap: int = 50) -> list[str]:
        """
        Splits the input text into chunks of specified size with optional overlap.

        :param text: The input text to be chunked.
        :param chunksize: The maximum size of each chunk.
        :param overlap: The number of characters to overlap between chunks.
        :return: A list of text chunks.
        """
        if chunksize <= 0:
            raise ValueError("chunksize must be a positive integer.")
        if overlap < 0:
            raise ValueError("overlap must be a non-negative integer.")

        chunks = []
        start = 0
        text_length = len(text)

        while start < text_length:
            end = min(start + chunksize, text_length)
            chunk = text[start:end]
            chunks.append(chunk)
            start += (chunksize - overlap)

        return chunks


